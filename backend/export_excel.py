"""Export the CR Portal database to a formatted Excel workbook.

Run any time to snapshot the live DB — including rows you entered through the
portal — to CR_Portal_Mock_Data.xlsx in the project root:

    cd backend
    .venv\\Scripts\\python export_excel.py
"""
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DB = HERE / "cr_portal.db"
OUT = HERE.parent / "CR_Portal_Mock_Data.xlsx"

NAVY = "1B2A44"
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(name="Arial", size=10)

conn = sqlite3.connect(str(DB))
cur = conn.cursor()

wb = Workbook()


def write_sheet(ws, title, headers, rows, num_formats=None):
    ws.title = title
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(list(row))
    for r in range(2, len(rows) + 2):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            if num_formats and headers[c - 1] in num_formats:
                cell.number_format = num_formats[headers[c - 1]]
    # column widths from content
    for c, h in enumerate(headers, start=1):
        width = max(
            len(str(h)),
            *(len(str(row[c - 1])) if row[c - 1] is not None else 1 for row in rows[:200]),
        ) if rows else len(str(h))
        ws.column_dimensions[get_column_letter(c)].width = min(width + 3, 28)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---- Sheet 1: Merchants (Table 1) ----
m_headers = [
    "MerchantID", "MerchantName", "Breadcrumb1", "Breadcrumb1Name",
    "Breadcrumb2", "Breadcrumb2Name", "AffiliateID", "AffiliateName",
    "Reporting", "Payout", "DealType", "RevenueStatus", "Owner",
]
cur.execute(
    """SELECT merchant_id, merchant_name, breadcrumb1, breadcrumb1_name,
              breadcrumb2, breadcrumb2_name, affiliate_id, affiliate_name,
              reporting, payout, deal_type, revenue_status, owner
       FROM merchants ORDER BY merchant_name"""
)
merchants = cur.fetchall()
write_sheet(wb.active, "Merchants", m_headers, merchants)

# ---- Sheet 2: Entries (Table 2) ----
e_headers = [
    "EntryID", "MerchantID", "MerchantName", "Date", "Clicks", "Sales",
    "CR %", "GMV", "Revenue", "Remarks", "RevenueStatus", "EnteredBy", "CreatedAt",
]
cur.execute(
    """SELECT e.id, e.merchant_id, m.merchant_name, e.entry_date, e.clicks,
              e.sales, e.cr, e.gmv, e.revenue, e.remarks, e.revenue_status,
              e.entered_by, e.created_at
       FROM entries e JOIN merchants m ON m.merchant_id = e.merchant_id
       ORDER BY e.entry_date DESC, m.merchant_name"""
)
entries = cur.fetchall()
ws2 = wb.create_sheet()
write_sheet(
    ws2, "Entries", e_headers, entries,
    num_formats={
        "Clicks": "#,##0", "Sales": "#,##0", "CR %": "0.00",
        "GMV": '"₹"#,##0.00', "Revenue": '"₹"#,##0.00',
    },
)

conn.close()
wb.save(OUT)
print(f"Wrote {OUT}: {len(merchants)} merchants, {len(entries)} entries")
